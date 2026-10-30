import json
import os
import pytest
from edocs.models import Document
from edocs.services.eimzo import EImzoService
from edocs.services.engine import CustomDocxTemplate, DocumentEngine
from edocs.services.qr import QRCodeService


def test_custom_docx_template_xml_patching():
    tpl = CustomDocxTemplate.__new__(CustomDocxTemplate)

    # 1. Test dollar-brace conversion to jinja2
    sample_xml = "<w:p><w:t>Shartnoma raqami: ${doc_number}</w:t></w:p>"
    patched = tpl.patch_xml(sample_xml)
    assert "{{doc_number}}" in patched
    assert "${doc_number}" not in patched

    # 2. Test dynamic loops conversion
    loop_xml = "<w:p><w:t>${for item in items}</w:t></w:p>"
    patched_loop = tpl.patch_xml(loop_xml)
    assert "{%for item in items" in patched_loop

    # 3. Test signers placeholder left alignment and RichText conversion
    signers_xml = '<w:p><w:jc w:val="right"/><w:t>${signers}</w:t></w:p>'
    patched_signers = tpl.patch_xml(signers_xml)
    assert "signers" in patched_signers
    assert '<w:jc w:val="left"/>' in patched_signers


def test_qr_code_generation():
    qr_path = QRCodeService.generate_document_qr(
        guid="test-guid-1234",
        view_url="https://test.uz/verify/test-guid-1234",
        security_code="9876",
    )
    assert os.path.exists(qr_path)
    assert qr_path.endswith(".png")
    assert os.path.getsize(qr_path) > 500  # Valid PNG image


@pytest.mark.django_db
def test_build_signing_payload():
    doc = Document.objects.create(
        name="Test Payload Doc",
        code="contract_grain",
        security_code="5566",
        extra_data={"contract_sum": "50000000"},
    )
    canonical_json, base64_payload = DocumentEngine.build_signing_payload(doc)

    assert canonical_json is not None
    assert base64_payload is not None

    parsed = json.loads(canonical_json)
    assert parsed["guid"] == str(doc.guid)
    assert parsed["security_code"] == "5566"
    assert parsed["data"]["contract_sum"] == "50000000"


def test_eimzo_service_fallback_verification():
    res = EImzoService.verify_pkcs7(
        pkcs7="DUMMY_BASE64_PKCS7_SIGNATURE",
        document_hash="test_sha256_hash",
        certificate_info={"CN": "Test Foydalanuvchi", "TIN": "301234567"},
    )
    assert res["status"] == 1
    assert "subjectCertificateInfo" in res
    assert res["subjectCertificateInfo"]["subjectName"]["CN"] == "Test Foydalanuvchi"


def test_patch_table_row_loops_list_context():
    tpl = CustomDocxTemplate.__new__(CustomDocxTemplate)
    tpl.current_context = {"items": [{"name": "A", "price": 10}, {"name": "B", "price": 20}]}

    table_xml = '<w:tbl><w:tr><w:tc><w:p><w:t>Nomi</w:t></w:p></w:tc></w:tr><w:tr><w:tc><w:p><w:t>{{items.name}} - {{items.price}}</w:t></w:p></w:tc></w:tr></w:tbl>'
    patched = tpl.patch_table_row_loops(table_xml)

    assert "{% for _row in (items or []) %}" in patched
    assert "{% endfor %}" in patched
    assert "_row.name" in patched
    assert "_row.price" in patched
    assert "{%tr" not in patched  # Must not use invalid {%tr %}


def test_patch_table_row_loops_skips_scalar_context():
    tpl = CustomDocxTemplate.__new__(CustomDocxTemplate)
    tpl.current_context = {"user": {"name": "Ali", "phone": "123"}, "contract_num": 42}

    table_xml = '<w:tbl><w:tr><w:tc><w:p><w:t>{{user.name}}</w:t></w:p></w:tc></w:tr></w:tbl>'
    patched = tpl.patch_table_row_loops(table_xml)
    assert "{% for" not in patched


def test_patch_table_row_loops_skips_existing_loops():
    tpl = CustomDocxTemplate.__new__(CustomDocxTemplate)
    tpl.current_context = {"items": [{"name": "A"}]}

    table_xml = '<w:tbl><w:tr><w:tc><w:p><w:t>{% for x in items %}{{x.name}}{% endfor %}</w:t></w:p></w:tc></w:tr></w:tbl>'
    patched = tpl.patch_table_row_loops(table_xml)
    assert patched.count("{% for") == 1


@pytest.mark.django_db
def test_generate_document_atomic_rollback_on_failure():
    import io
    from unittest.mock import patch
    from django.contrib.auth import get_user_model
    from django.core.files.uploadedfile import SimpleUploadedFile
    from docx import Document as DocxDocument
    from edocs.exceptions import DocumentGenerationError
    from edocs.models import DocumentTemplate
    from tests.models import SampleContract

    user = get_user_model().objects.create_user(username="atomic_user", password="pwd")
    contract = SampleContract.objects.create(title="Atomic Test Contract", creator=user, status="new")

    docx_doc = DocxDocument()
    docx_doc.add_paragraph("Hello world")
    bio = io.BytesIO()
    docx_doc.save(bio)
    bio.seek(0)

    DocumentTemplate.objects.create(
        name="Fail Template",
        code="fail_tpl",
        file=SimpleUploadedFile("fail.docx", bio.read(), content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    )

    initial_docs_count = Document.objects.count()

    with patch.object(CustomDocxTemplate, "render", side_effect=RuntimeError("Docx render crash")):
        with pytest.raises(DocumentGenerationError):
            DocumentEngine.generate_document(
                template_code="fail_tpl",
                context_data={},
                related_object=contract,
            )

    # Verify no orphan Document row was left in DB
    assert Document.objects.count() == initial_docs_count
    contract.refresh_from_db()
    assert contract.status == "new"
    assert contract.document is None


@pytest.mark.django_db
def test_generate_document_updates_host_status_to_pending():
    import io
    from unittest.mock import patch
    from django.contrib.auth import get_user_model
    from django.core.files.uploadedfile import SimpleUploadedFile
    from docx import Document as DocxDocument
    from edocs.models import DocumentTemplate
    from tests.models import SampleContract

    user = get_user_model().objects.create_user(username="pending_user", password="pwd")
    contract = SampleContract.objects.create(title="Pending Test Contract", creator=user, status="new")

    docx_doc = DocxDocument()
    docx_doc.add_paragraph("Doc: {{doc_number}}")
    bio = io.BytesIO()
    docx_doc.save(bio)
    bio.seek(0)

    DocumentTemplate.objects.create(
        name="Pending Template",
        code="pending_tpl",
        file=SimpleUploadedFile("pending.docx", bio.read(), content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    )

    with patch.object(DocumentEngine, "convert_to_pdf", return_value=None):
        doc = DocumentEngine.generate_document(
            template_code="pending_tpl",
            context_data={"doc_number": "1001"},
            related_object=contract,
        )

    contract.refresh_from_db()
    assert contract.status == "pending"
    assert contract.document == doc


def test_dependency_floors_for_cve_fixes():
    import tomllib
    from pathlib import Path

    pyproject_path = Path(__file__).resolve().parent.parent / "pyproject.toml"
    with open(pyproject_path, "rb") as f:
        data = tomllib.load(f)

    deps = {d.split(">=")[0]: d.split(">=")[1] for d in data["project"]["dependencies"] if ">=" in d}
    assert deps["Django"] >= "4.2.18"
    assert deps["djangorestframework"] >= "3.15.2"
    assert deps["jinja2"] >= "3.1.6"
    assert deps["requests"] >= "2.32.3"
    assert deps["pillow"] >= "10.4.0"
